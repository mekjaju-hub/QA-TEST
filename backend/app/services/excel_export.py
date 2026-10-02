"""Excel Export 10 Sheets (หัวข้อ 37) — openpyxl.

Fixes prototype Known Issue: Freeze Header + Status colours now supported.
Thai text, wrap text, autofilter, column width, long code kept intact (cell limit 32,767 chars → split notice).
Copy saved to storage/exports/.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    AutomationArtifact, Project, Requirement, RequirementConflict, TestCase, TestRun, TestRunResult, TestScenario,
)
from .storage import get_storage

STATUS_FILL = {
    "APPROVED": "E4F5EC", "PASSED": "E4F5EC", "READY_FOR_AUTOMATION": "E4F5EC", "AUTOMATED": "E4F5EC", "RESOLVED": "E4F5EC",
    "NEEDS_CLARIFICATION": "FDF1DF", "WAITING_FOR_REVIEW": "FDF1DF", "REVISED": "FDF1DF", "BLOCKED": "FDF1DF",
    "CONFLICT": "FCE8E6", "FAILED": "FCE8E6", "OPEN": "FCE8E6",
    "AI_GENERATED": "F0E9FB", "DRAFT": "EEF1F4", "DEPRECATED": "EEF1F4",
}
HEADER_FILL = PatternFill("solid", fgColor="13233A")
MAX_CELL = 32000


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        v = "\n".join(str(x) for x in v)
    if isinstance(v, str) and len(v) > MAX_CELL:
        return v[:MAX_CELL] + "\n…[ตัดที่ 32,000 ตัวอักษร — ดูไฟล์เต็มใน Automation ZIP]"
    return v


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int], status_col: str | None = "Status") -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([_cell(v) for v in r])
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    if status_col and status_col in header:
        idx = header.index(status_col) + 1
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                col = STATUS_FILL.get(str(c.value))
                if col:
                    c.fill = PatternFill("solid", fgColor=col)


def build_workbook(db: Session, pr: Project, username: str) -> Workbook:
    reqs = list(db.scalars(select(Requirement).where(Requirement.project_id == pr.id).order_by(Requirement.req_id)))
    rq = {r.id: r for r in reqs}
    scs = list(db.scalars(select(TestScenario).where(TestScenario.project_id == pr.id).order_by(TestScenario.ts_id)))
    sc = {s.id: s for s in scs}
    tcs = list(db.scalars(select(TestCase).where(TestCase.project_id == pr.id).order_by(TestCase.tc_id)))
    confs = list(db.scalars(select(RequirementConflict).where(RequirementConflict.project_id == pr.id)))
    arts = list(db.scalars(select(AutomationArtifact).where(AutomationArtifact.project_id == pr.id)))
    latest = [r for r in reqs if r.is_latest]
    wb = Workbook()
    wb.remove(wb.active)
    approved = sum(1 for t in tcs if t.status in ("APPROVED", "READY_FOR_AUTOMATION", "AUTOMATED"))
    _sheet(wb, "Summary", ["Item", "Value"], [
        ["Project", f"{pr.code} - {pr.name}"], ["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")], ["Exported by", username],
        ["Requirements", len(latest)], ["Needs Clarification", sum(1 for r in latest if r.status == "NEEDS_CLARIFICATION")],
        ["Conflicts (open)", sum(1 for c in confs if c.status == "OPEN")], ["Test Scenarios", len(scs)], ["Test Cases", len(tcs)],
        ["Approved", approved]], [26, 50], status_col=None)
    _sheet(wb, "Requirements", ["Requirement ID", "Version", "Type", "Module", "Title", "Original Text", "Expected Result", "Role",
                                "Threshold", "Unit", "Date Range", "Inclusion", "Exclusion", "Source Doc", "Doc Version", "Source Page",
                                "Source Section", "Completeness", "Completeness Reasons", "Clarity", "Clarity Reasons", "Status"],
           [[r.req_id, r.version, r.type, r.module, r.title, r.original_text, r.expected_result, r.role, r.threshold, r.unit, r.date_range,
             r.inclusion, r.exclusion, r.source.document_name if r.source else "", r.doc_version, r.source.page if r.source else "",
             r.source.section if r.source else "", r.completeness_score,
             [("✓ " if x["ok"] else "✗ ") + x["text"] for x in r.completeness_reasons], r.clarity_score,
             [("✓ " if x["ok"] else "✗ ") + x["text"] for x in r.clarity_reasons], r.status] for r in reqs],
           [24, 8, 18, 12, 36, 60, 36, 18, 14, 8, 18, 24, 24, 20, 8, 10, 24, 12, 40, 8, 40, 20])
    _sheet(wb, "Clarification Questions", ["Requirement ID", "Question", "Answer", "Resolved", "Assumption", "Assumption State", "Source Page", "Status"],
           [[r.req_id, q.text, q.answer, "Yes" if q.resolved else "No",
             ("AI ASSUMPTION - NOT FOUND IN BRS: " + q.assumption.text) if q.assumption else "", q.assumption.state if q.assumption else "",
             r.source.page if r.source else "", "RESOLVED" if q.resolved else "OPEN"] for r in latest for q in r.questions],
           [24, 50, 40, 10, 50, 14, 10, 12])
    _sheet(wb, "Conflicts", ["Requirement A", "Requirement B", "Differences", "Status", "Resolution", "Reason"],
           [[rq[c.req_a_id].req_id if c.req_a_id in rq else "", rq[c.req_b_id].req_id if c.req_b_id in rq else "",
             [f"{d['field']}: {d['a']} ≠ {d['b']}" for d in c.diffs], c.status, c.resolution, c.reason] for c in confs],
           [24, 24, 50, 12, 30, 40])
    _sheet(wb, "Test Scenarios", ["Scenario ID", "Requirement ID", "Title", "Type", "Priority", "Risk", "Rationale", "Status"],
           [[s.ts_id, rq[s.requirement_id].req_id if s.requirement_id in rq else "", s.title, s.type, s.priority, s.risk,
             s.rationale, s.status] for s in scs], [24, 24, 50, 12, 10, 8, 50, 18])
    _sheet(wb, "Test Cases", ["Test Case ID", "Version", "Scenario ID", "Requirement ID", "Requirement Version", "Title", "Business Explanation",
                              "Given", "When", "Then", "Priority", "Risk", "Priority/Risk Reasons", "Origin", "Assumption", "Clarification",
                              "Source Page", "Status", "Approved By", "Approved Date"],
           [[t.tc_id, t.version, sc[t.scenario_id].ts_id if t.scenario_id in sc else "", rq[t.requirement_id].req_id if t.requirement_id in rq else "",
             t.req_version, t.title, t.business_explanation, t.given, t.when, t.then, t.priority, t.risk, t.pr_reasons, t.origin,
             t.assumption, t.clarification_ref, t.source_page, t.status, t.approved_by or "",
             t.approved_at.strftime("%Y-%m-%d %H:%M") if t.approved_at else ""] for t in tcs],
           [24, 8, 24, 24, 10, 40, 50, 40, 40, 40, 10, 8, 40, 18, 40, 40, 10, 20, 14, 16])
    _sheet(wb, "Test Steps", ["Test Case ID", "Version", "Step", "Action", "Test Data", "Expected Result", "Origin/Label"],
           [[t.tc_id, t.version, s.step_no, s.action, s.test_data, s.expected, " | ".join(x for x in (s.origin, s.label) if x)]
            for t in tcs for s in t.steps], [24, 8, 6, 50, 30, 40, 40], status_col=None)
    _sheet(wb, "Test Data", ["Test Case ID", "Value", "Expected", "Origin", "Label"],
           [[t.tc_id, d.value, d.expected, d.origin, d.label] for t in tcs for d in t.data_rows], [24, 24, 30, 20, 50], status_col=None)
    trace_rows = []
    for t in tcs:
        r = rq.get(t.requirement_id)
        res = db.execute(select(TestRunResult.status).join(TestRun, TestRun.id == TestRunResult.run_id)
                         .where(TestRunResult.test_case_id == t.id).order_by(TestRun.started_at.desc()).limit(1)).first()
        trace_rows.append([r.source.document_name if r and r.source else "", r.doc_version if r else "", r.source.section if r and r.source else "",
                           r.req_id if r else "", sc[t.scenario_id].ts_id if t.scenario_id in sc else "", t.tc_id,
                           [f"{a.kind}:{a.name}" for a in arts if t.id in (a.test_case_ids or [])], res[0] if res else ""])
    _sheet(wb, "Traceability", ["Document", "Version", "Section", "Requirement ID", "Scenario ID", "Test Case ID", "Artifacts", "Status"],
           trace_rows, [24, 8, 30, 24, 24, 24, 40, 16])
    auto_rows = []
    for t in tcs:
        mine = [a for a in arts if t.id in (a.test_case_ids or [])]
        code = ""
        py = next((a for a in mine if a.kind == "pytest"), None)
        if py:
            st = get_storage()
            f = next((x for x in py.files if x.startswith("tests/unit/")), None)
            if f and st.exists(f"{py.storage_path}/current/{f}"):
                code = st.read_text(f"{py.storage_path}/current/{f}")
        auto_rows.append([t.tc_id, t.status, t.automation_candidate, t.automation_tool, ", ".join(a.kind for a in mine), code])
    _sheet(wb, "Automation Status", ["Test Case ID", "Status", "Automation Candidate", "Tool", "Artifacts", "Code (first test file)"],
           auto_rows, [24, 20, 12, 18, 24, 100])
    return wb


def export_project_xlsx(db: Session, pr: Project, username: str) -> tuple[bytes, str]:
    wb = build_workbook(db, pr, username)
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    key = f"exports/{pr.code}_QA_TestCases_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    get_storage().write_bytes(key, data)
    return data, key
