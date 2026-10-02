"""Document parsers (หัวข้อ 5, 6 Step 1–2) — Python port of parseDocx/parseXlsx/parseCsv/parsePdf/parseTxt.

Every parser returns {"blocks": [...], "images": [...], "warnings": [...], "page_count": int|None, "encoding": str}
Block: {"type": "heading"|"para"|"table"|"image", "text", "rows", "caption", "comments", "page", "image_id"}
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from collections import Counter

from ...core.errors import AppError
from .text import is_heading, normalize_text, text_to_blocks

ALLOWED_EXT = ["docx", "xlsx", "csv", "pdf", "txt"]
CAPTION_RE = re.compile(r"^(รูป|ภาพ|Figure|Screen|หน้าจอ)", re.I)


def ext_of(name: str) -> str:
    return (str(name).rsplit(".", 1)[-1] if "." in str(name) else "").lower()


def sniff_type(data: bytes, ext: str) -> bool:
    """Validate file content against extension (magic bytes)."""
    head = data[:8]
    is_zip = head[:2] == b"PK"
    is_pdf = head[:4] == b"%PDF"
    if ext in ("docx", "xlsx"):
        return is_zip
    if ext == "pdf":
        return is_pdf
    if ext in ("csv", "txt", "paste"):
        if data.startswith(b"\xff\xfe"):
            return True
        return not is_zip and not is_pdf and b"\x00" not in head
    return False


def decode_text(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace"), "UTF-8 (BOM)"
    if data.startswith(b"\xff\xfe"):
        return data[2:].decode("utf-16-le", errors="replace"), "UTF-16LE"
    try:
        return data.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        return data.decode("cp874", errors="replace"), "Windows-874 (TIS-620)"


# ------------------------------------------------------------------ DOCX
def parse_docx(data: bytes) -> dict:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise AppError("PARSER_MISSING", "ไม่พบ Library python-docx", technical=str(exc)) from exc
    try:
        doc = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise AppError("CORRUPT_FILE", "ไฟล์ DOCX เสียหายหรืออ่านไม่ได้", technical=str(exc), action="ตรวจไฟล์แล้วอัปโหลดใหม่") from exc

    blocks: list[dict] = []
    images: list[dict] = []
    warnings: list[str] = []
    last_img: dict | None = None
    ns_blip = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    ns_rid = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            # embedded images (screenshots)
            for blip in child.iter(ns_blip):
                rid = blip.get(ns_rid)
                part = doc.part.related_parts.get(rid) if rid else None
                if part is None:
                    continue
                img = {"id": str(uuid.uuid4()), "content_type": getattr(part, "content_type", "image/png"),
                       "blob": part.blob, "alt": "", "caption": "",
                       "near_text": blocks[-1]["text"][:200] if blocks and blocks[-1].get("text") else "",
                       "status": "NEEDS_VISUAL_REVIEW"}
                images.append(img)
                blocks.append({"type": "image", "image_id": img["id"], "page": None})
                last_img = img
            txt = normalize_text(p.text)
            if not txt:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if last_img and (CAPTION_RE.match(txt) or style.lower() == "caption"):
                last_img["caption"] = txt
                last_img = None
            if style.lower().startswith(("heading", "title")):
                blocks.append({"type": "heading", "text": txt, "page": None})
            elif "list" in style.lower():
                blocks.append({"type": "para", "text": "- " + txt, "page": None})
            else:
                heading = is_heading(txt) and len(txt) < 90 and not re.search(r"ต้อง|shall|must|ระบบ", txt, re.I)
                blocks.append({"type": "heading" if heading else "para", "text": txt, "page": None})
        elif tag == "tbl":
            t = Table(child, doc)
            rows = []
            for r in t.rows:
                cells = []
                prev = None
                for c in r.cells:  # merged cells repeat the same _tc → keep once
                    if prev is not None and c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(normalize_text(c.text))
                rows.append(cells)
            rows = [r for r in rows if any(r)]
            if rows:
                blocks.append({"type": "table", "rows": rows, "page": None})
    # Word comments (python-docx >= 1.2)
    try:
        for cm in getattr(doc, "comments", []) or []:
            txt = normalize_text(getattr(cm, "text", ""))
            if txt:
                blocks.append({"type": "para", "text": f"หมายเหตุ (Comment): {txt}", "page": None})
    except Exception:  # noqa: BLE001
        warnings.append("อ่าน Comment ใน DOCX ไม่ได้")
    warnings.append("DOCX ไม่มีเลขหน้าในตัว — Source Page = NOT_FOUND (ใช้ Section แทน)")
    return {"blocks": blocks, "images": images, "warnings": warnings, "page_count": None, "encoding": "UTF-8 (OOXML)"}


# ------------------------------------------------------------------ XLSX
def parse_xlsx(data: bytes) -> dict:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise AppError("CORRUPT_FILE", "ไฟล์ XLSX เสียหายหรืออ่านไม่ได้", technical=str(exc)) from exc
    blocks = []
    for ws in wb.worksheets:
        rows, comments = [], []
        for row in ws.iter_rows():
            vals = []
            for c in row:
                vals.append(normalize_text("" if c.value is None else c.value))
                if c.comment is not None:
                    comments.append({"cell": c.coordinate, "text": normalize_text(c.comment.text)})
            rows.append(vals)
        # trim trailing empty columns
        width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
        rows = [r[:width] for r in rows if any(r)]
        merges = len(ws.merged_cells.ranges)
        if rows:
            cap = f"Worksheet: {ws.title}" + (f" (merged cells: {merges})" if merges else "")
            blocks.append({"type": "table", "rows": rows, "caption": cap, "comments": comments, "page": None})
            for cm in comments:
                blocks.append({"type": "para", "text": f"หมายเหตุ Cell {cm['cell']}: {cm['text']}", "page": None})
    return {"blocks": blocks, "images": [], "warnings": [f"{len(wb.worksheets)} worksheet"], "page_count": None, "encoding": "UTF-8 (OOXML)"}


# ------------------------------------------------------------------ CSV
def parse_csv(data: bytes) -> dict:
    text, encoding = decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    rows = [[normalize_text(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise AppError("EMPTY_FILE", "ไฟล์ CSV ไม่มีข้อมูล")
    return {"blocks": [{"type": "table", "rows": rows, "caption": "CSV", "page": None}], "images": [],
            "warnings": [f"Encoding: {encoding}", f'Delimiter: "{delim}"'], "page_count": None, "encoding": encoding}


# ------------------------------------------------------------------ PDF
PAGE_NO = re.compile(r"^(?:page|หน้า|p\.)?\s*\d+(?:\s*(?:of|/|จาก)\s*\d+)?$", re.I)


def parse_pdf(data: bytes) -> dict:
    try:
        import pdfplumber
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise AppError("CORRUPT_FILE", "ไฟล์ PDF เสียหายหรืออ่านไม่ได้", technical=str(exc)) from exc
    pages: list[list[str]] = []
    with pdf:
        for pg in pdf.pages:
            txt = pg.extract_text() or ""
            pages.append([ln for ln in (normalize_text(x) for x in txt.split("\n")) if ln])
    n = len(pages)
    key_of = lambda ln: "#PAGE#" if PAGE_NO.match(ln) else ln  # noqa: E731
    freq: Counter = Counter()
    for ls in pages:
        freq.update({key_of(x) for x in ls[:2] + ls[-2:]})
    removed = {k for k, v in freq.items() if n >= 3 and v / n > 0.5}
    blocks = []
    for i, ls in enumerate(pages):
        for idx, ln in enumerate(ls):
            if (idx < 2 or idx >= len(ls) - 2) and key_of(ln) in removed:
                continue
            blocks.append({"type": "heading" if is_heading(ln) else "para", "text": ln, "page": i + 1})
    empty = sum(1 for ls in pages if not ls)
    if n == 0 or empty > n * 0.5:
        raise AppError("PDF_NO_TEXT", "PDF นี้ไม่มีข้อความที่เลือกได้ (อาจเป็น PDF Scan)", technical=f"empty pages: {empty}/{n}",
                       action="Version นี้ยังไม่รองรับ OCR — ใช้ไฟล์ DOCX หรือ PDF ที่เลือกข้อความได้")
    warnings = [f"{n} หน้า"] + ([f"ลบ Header/Footer ซ้ำ {len(removed)} รูปแบบ"] if removed else [])
    return {"blocks": blocks, "images": [], "warnings": warnings, "page_count": n, "encoding": "PDF text layer"}


# ------------------------------------------------------------------ TXT / Paste
def parse_txt(data: bytes) -> dict:
    text, encoding = decode_text(data)
    return {"blocks": text_to_blocks(text), "images": [], "warnings": [f"Encoding: {encoding}"], "page_count": None, "encoding": encoding}


def parse_document(data: bytes, ext: str) -> dict:
    ext = ext.lower()
    if ext == "docx":
        return parse_docx(data)
    if ext == "xlsx":
        return parse_xlsx(data)
    if ext == "csv":
        return parse_csv(data)
    if ext == "pdf":
        return parse_pdf(data)
    if ext in ("txt", "paste"):
        return parse_txt(data)
    raise AppError("UNSUPPORTED_FILE", "ไม่รองรับไฟล์ประเภทนี้", technical=ext, action="ใช้ DOCX, XLSX, CSV, PDF หรือ TXT")
