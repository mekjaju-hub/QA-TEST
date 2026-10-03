"""Test Case Library: every test case from every page history, merged into one list without duplicates.

Two test cases are "the same" when they are in the same category and have the same title after normalising
(case, spaces, quote styles). The merged row keeps references to every website/page it came from (page + WP-ID),
so you can still trace it back. Library IDs (TL-001 …) are stable and stored in storage/explore/library_index.json.
"""
from __future__ import annotations

import io
import json
import re
import threading

from .storage import get_storage
from .web_history import HBASE

INDEX = "explore/library_index.json"

CATEGORIES = {  # sig prefix → (category code, Thai label)
    "login": ("login", "Login"),
    "home": ("after_login", "หลัง Login"),
    "click": ("click", "กดปุ่ม/ลิงก์"),
    "field": ("input", "กรอกข้อมูล"),
    "dropdown": ("dropdown", "Dropdown"),
    "page": ("page", "หน้าเว็บ"),
    "record": ("record", "สถานการณ์ที่บันทึก"),
    "record-neg": ("record", "สถานการณ์ที่บันทึก"),
}
CATEGORY_LABEL = {code: label for code, label in CATEGORIES.values()}
URL_RX = re.compile(r"https?://\S+")


def category_of(sig: str) -> str:
    head = re.split(r"[.:]", sig, maxsplit=1)[0]
    return CATEGORIES.get(head, ("other", "อื่นๆ"))[0]


def norm(text: str) -> str:
    t = (text or "").lower().replace("“", '"').replace("”", '"').replace("'", '"')
    return re.sub(r"\s+", " ", t).strip()


def generic_steps(steps: list[str]) -> list[str]:
    return [URL_RX.sub("[หน้าเว็บที่ทดสอบ]", s) for s in steps]


def _histories() -> list[dict]:
    st = get_storage()
    out = []
    for k in st.list(HBASE):
        if k.endswith(".json"):
            try:
                out.append(json.loads(st.read_text(f"{HBASE}/{k}")))
            except Exception:  # noqa: BLE001
                continue
    return out


_build_lock = threading.Lock()


def build() -> list[dict]:
    with _build_lock:          # one request at a time assigns new TL-IDs and writes the index
        return _build()


def _build() -> list[dict]:
    st = get_storage()
    index: dict[str, str] = json.loads(st.read_text(INDEX)) if st.exists(INDEX) else {}
    known = len(index)
    next_no = max([int(v.split("-")[1]) for v in index.values()] + [0]) + 1
    rows: dict[str, dict] = {}
    for h in sorted(_histories(), key=lambda x: x.get("created_at", "")):
        for sig, t in sorted(h["test_cases"].items(), key=lambda kv: kv[1].get("hid", "")):
            if "first_seen" not in t:
                continue
            cat = category_of(sig)
            key = f"{cat}|{norm(t['title'])}"
            if key not in index:
                index[key] = f"TL-{next_no:03d}"
                next_no += 1
            row = rows.setdefault(key, {
                "lid": index[key], "key": key, "category": cat, "category_label": CATEGORY_LABEL.get(cat, "อื่นๆ"),
                "title": t["title"], "type": t.get("type", ""), "priority": t.get("priority", ""),
                "steps": generic_steps(t.get("steps", [])), "expected": URL_RX.sub("[หน้าเว็บที่ทดสอบ]", t.get("expected", "")),
                "needs_login": bool(t.get("needs_login")), "first_seen": t["first_seen"], "last_seen": t["last_seen"],
                "sources": [], "runs": 0, "passed": 0})
            row["first_seen"] = min(row["first_seen"], t["first_seen"])
            row["last_seen"] = max(row["last_seen"], t["last_seen"])
            row["runs"] += t.get("runs", 0)
            row["passed"] += t.get("passed", 0)
            if {"high": 3, "medium": 2, "low": 1}.get((t.get("priority") or "").lower(), 0) > \
               {"high": 3, "medium": 2, "low": 1}.get(row["priority"].lower(), 0):
                row["priority"] = t["priority"]
            row["sources"].append({"page_key": h["key"], "page": h["page"], "title": h.get("title", ""), "hid": t["hid"],
                                   "last_result": t.get("last_result"), "last_run_at": t.get("last_run_at")})
    if len(index) != known:    # only when new TL-IDs were given (reading the library must not rewrite files)
        st.write_text(INDEX, json.dumps(index, ensure_ascii=False, indent=1))
    out = []
    for r in rows.values():
        ran = [s for s in r["sources"] if s["last_result"]]
        latest = max(ran, key=lambda s: s["last_run_at"] or "", default=None)
        r["last_result"] = latest["last_result"] if latest else None
        r["sites"] = len({s["page"].split("/")[0] for s in r["sources"]})
        r["pages"] = len(r["sources"])
        out.append(r)
    return sorted(out, key=lambda r: r["lid"])


def filtered(rows: list[dict], q: str = "", category: str = "", result: str = "", site: str = "") -> list[dict]:
    qn = norm(q)
    out = []
    for r in rows:
        if category and r["category"] != category:
            continue
        if result == "none" and r["last_result"]:
            continue
        if result and result != "none" and r["last_result"] != result:
            continue
        if site and not any(site.lower() in s["page"].lower() for s in r["sources"]):
            continue
        if qn and qn not in norm(" ".join([r["lid"], r["title"], r["expected"], " ".join(r["steps"]),
                                             " ".join(s["page"] + " " + s["hid"] for s in r["sources"])])):
            continue
        out.append(r)
    return out


HEADERS = ["Library ID", "หมวด", "Test Case", "ประเภท", "Priority", "ขั้นตอน", "ผลที่คาดหวัง", "ต้อง Login",
           "จำนวนเว็บไซต์", "พบในหน้า (WP-ID)", "ผลรันล่าสุด", "รัน/ผ่าน", "พบครั้งแรก", "ล่าสุด"]


def table(rows: list[dict]) -> list[list[str]]:
    data = []
    for r in rows:
        data.append([r["lid"], r["category_label"], r["title"], r["type"], r["priority"],
                     "\n".join(f"{i + 1}. {s}" for i, s in enumerate(r["steps"])), r["expected"], "ใช่" if r["needs_login"] else "",
                     str(r["sites"]), "\n".join(f"{s['page']} ({s['hid']})" for s in r["sources"]), r["last_result"] or "",
                     f"{r['runs']}/{r['passed']}", r["first_seen"][:10], r["last_seen"][:10]])
    return data


def to_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Case Library"
    ws.append(HEADERS)
    for line in table(rows):
        ws.append(line)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B2A44")
    widths = [11, 14, 42, 13, 9, 48, 40, 9, 9, 36, 12, 9, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        res = row[10]
        if res.value == "PASSED":
            res.fill = PatternFill("solid", fgColor="DDF3E4")
        elif res.value == "FAILED":
            res.fill = PatternFill("solid", fgColor="FBE2E0")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
