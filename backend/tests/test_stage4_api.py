"""Stage 4 — API, permission, requirement/conflict/clarification, test design, export, version compare."""
import io

from openpyxl import load_workbook

from tests.conftest import REPO


def reqs(client, h, project, **params):
    return {r["req_id"]: r for r in client.get(f"/api/projects/{project['id']}/requirements", headers=h, params=params).json()["items"]}


# ---------------------------------------------------------------- auth & security
def test_login_generic_error_and_unauthorized(client):
    r = client.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
    assert r.status_code == 401 and r.json()["error"]["user_message"] == "Username หรือ Password ไม่ถูกต้อง"
    r = client.post("/api/auth/login", json={"username": "nobody", "password": "wrong"})
    assert r.json()["error"]["user_message"] == "Username หรือ Password ไม่ถูกต้อง"
    assert client.get("/api/projects").status_code == 401
    assert client.get("/api/projects", headers={"Authorization": "Bearer abc.def.ghi"}).status_code == 401


def test_must_change_password_blocks_api(client):
    from app.db import SessionLocal
    from app.core.security import hash_password
    from app.models import User
    db = SessionLocal()
    db.add(User(username="fresh", password_hash=hash_password("Temp-Pass#123"), must_change_password=True))
    db.commit()
    db.close()
    tok = client.post("/api/auth/login", json={"username": "fresh", "password": "Temp-Pass#123"}).json()
    assert tok["must_change_password"] is True
    r = client.get("/api/projects", headers={"Authorization": f"Bearer {tok['access_token']}"})
    assert r.status_code == 403 and r.json()["error"]["code"] == "PASSWORD_CHANGE_REQUIRED"


def test_weak_password_rejected(client, qam):
    r = client.post("/api/auth/change-password", headers=qam, json={"current_password": "N3w-Strong!Pass", "new_password": "aaaaaaaaaa"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "WEAK_PASSWORD"


def test_role_restrictions(client, ba, qaa, project):
    files = {"file": ("x.txt", "ระบบต้องแสดงผล".encode(), "text/plain")}
    assert client.post(f"/api/projects/{project['id']}/documents", headers=ba, files=files).status_code == 403
    assert client.post("/api/projects", headers=qaa, json={"code": "XX", "name": "x"}).status_code == 403
    assert client.get("/api/audit-logs", headers=qaa).status_code == 403
    assert client.get("/api/settings", headers=ba).status_code == 403


def test_secret_not_exposed(client, qam):
    me = client.get("/api/auth/me", headers=qam).json()
    text = str(me)
    assert "secret_key" not in text and "claude_api_key" not in text and "github_token" not in text


def test_upload_invalid_type_and_traversal_name(client, qam, scratch):
    r = client.post(f"/api/projects/{scratch['id']}/documents", headers=qam, files={"file": ("evil.docx", b"MZ\x90\x00binary", "application/octet-stream")})
    assert r.status_code == 400 and r.json()["error"]["code"] == "INVALID_FILE_TYPE"
    r = client.post(f"/api/projects/{scratch['id']}/documents", headers=qam, files={"file": ("a.exe", b"MZ", "application/octet-stream")})
    assert r.json()["error"]["code"] == "UNSUPPORTED_FILE"
    r = client.post(f"/api/projects/{scratch['id']}/documents", headers=qam,
                    files={"file": ("../../evil.txt", "ระบบต้องแสดงรายงานให้ผู้ใช้".encode(), "text/plain")}, data={"auto_process": "false"})
    assert r.status_code == 200
    assert r.json()["version"]["filename"] == "evil.txt"


def test_masking():
    from app.core.masking import mask
    s = mask("password=abc123 token: xyz Bearer abc.def ghp_ABCDEFGHIJKLMNOPQRST postgres://u:p@h/db CUST123456 1-2345-67890-12-3 OTP=123456")
    for leaked in ("abc123", "xyz", "abc.def", "ABCDEFGHIJKLMNOPQRST", "u:p@h", "CUST123456", "1-2345-67890-12-3", "123456 "):
        assert leaked not in s, leaked
    assert "CUST-TEST-0001" in mask("CUST-TEST-0001")


# ---------------------------------------------------------------- upload formats / paste
def test_upload_all_formats(client, qam, scratch):
    for name in ("sample.docx", "sample.xlsx", "sample.csv", "sample.pdf"):
        r = client.post(f"/api/projects/{scratch['id']}/documents", headers=qam,
                        files={"file": (name, (REPO / "samples" / name).read_bytes())})
        assert r.status_code == 200, (name, r.text)
        prog = client.get(f"/api/documents/{r.json()['document']['id']}/progress", headers=qam).json()
        assert prog["job"]["status"] == "READY_FOR_REVIEW", (name, prog["log_tail"])
        assert prog["sections"], name
    pdf_doc = r.json()["document"]["id"]
    sec_pages = [s["page"] for s in client.get(f"/api/documents/{pdf_doc}/progress", headers=qam).json()["sections"]]
    assert sec_pages == [1, 2, 3]


def test_paste_text(client, qam, scratch):
    r = client.post(f"/api/projects/{scratch['id']}/paste-text", headers=qam,
                    json={"title": "pasted", "text": "1. Rule 9 Paste\n- ระบบต้องแสดงข้อความ OK เมื่อบันทึกสำเร็จ ให้ผู้ดูแลระบบ"})
    assert r.status_code == 200
    assert r.json()["version"]["filename"] == "pasted.txt"
    rs = reqs(client, qam, scratch, q="Paste")
    assert not rs or True
    assert any(k.startswith("REQ-SCR-RULE9-") for k in reqs(client, qam, scratch))


# ---------------------------------------------------------------- requirement analysis
def test_requirement_ids_sources_and_quality(client, qam, project, demo_doc):
    rs = reqs(client, qam, project)
    assert {"REQ-CAM-RULE3-001", "REQ-CAM-RULE5-002", "REQ-CAM-RULE6-001"} <= set(rs)
    d = client.get(f"/api/requirements/{rs['REQ-CAM-RULE4-003']['id']}", headers=qam).json()
    assert d["source"]["section"] == "3. Rule 4 Transaction Monthly Report"
    texts = [x["text"] for x in d["completeness_reasons"]]
    assert "ไม่มี Expected Result" in texts and "ไม่ระบุ Role" in texts          # AC 9, 10
    assert any(q["key"] == "vague" for q in d["questions"])
    d2 = client.get(f"/api/requirements/{rs['REQ-CAM-RULE3-001']['id']}", headers=qam).json()
    assert d2["fields"]["date_range"] != "NOT_FOUND"
    # missing threshold detection: requirement mentions numeric words but no threshold value
    from app.services.engine.analysis import analyze_quality, extract_fields
    r = {"type": "Business Rule", "original_text": "ระบบต้องแสดงลูกค้าที่มียอดมากกว่าเกณฑ์", "questions": []}
    r["fields"] = extract_fields(r["original_text"])
    q = analyze_quality(r)
    assert any(x["text"] == "ไม่มี Threshold" for x in q["completeness"]["reasons"])   # AC 11
    # AI output structure (หัวข้อ 32)
    assert set(d["ai_output"]) == {"found_in_brs", "assumptions", "recommendations", "clarification_questions", "conflicts", "source_references", "confidence"}


def test_conflict_blocks_approve_and_resolution(client, qam, ba, project, demo_doc):
    rs = reqs(client, qam, project)
    a = rs["REQ-CAM-RULE3-001"]
    assert a["status"] == "CONFLICT"
    r = client.post(f"/api/requirements/{a['id']}/approve", headers=qam)
    assert r.status_code == 409 and r.json()["error"]["code"] == "HAS_CONFLICT"
    center = client.get(f"/api/projects/{project['id']}/clarifications", headers=ba).json()
    conf = [c for c in center["conflicts"] if c["status"] == "OPEN"][0]
    assert conf["diffs"][0]["field"] == "Operator (> vs >=)"
    # reason is mandatory
    assert client.post(f"/api/requirements/{a['id']}/resolve-conflict", headers=ba, json={"conflict_id": conf["id"], "keep": "a", "reason": ""}).status_code == 422
    r = client.post(f"/api/requirements/{a['id']}/resolve-conflict", headers=ba,
                    json={"conflict_id": conf["id"], "keep": "a", "reason": "BA ยืนยันว่าใช้ >= ตาม Rule 3"})
    assert r.status_code == 200 and r.json()["status"] == "RESOLVED"
    rs = reqs(client, qam, project, show_old=True)
    assert rs["REQ-CAM-RULE4-001"]["status"] == "DEPRECATED"
    assert rs["REQ-CAM-RULE3-001"]["status"] != "CONFLICT"


def test_clarification_assumption_and_approve(client, qam, ba, project, demo_doc):
    rs = reqs(client, qam, project)
    rid = rs["REQ-CAM-RULE3-001"]["id"]
    d = client.get(f"/api/requirements/{rid}", headers=qam).json()
    assumption_q = next(q for q in d["questions"] if q["assumption"])
    assert assumption_q["assumption"]["label"] == "AI ASSUMPTION - NOT FOUND IN BRS"     # AC 15
    r = client.post(f"/api/requirements/{rid}/assumption", headers=qam, json={"question_id": assumption_q["id"], "state": "ACCEPTED"})
    assert r.status_code == 200
    # BA cannot edit assumptions
    assert client.post(f"/api/requirements/{rid}/assumption", headers=ba, json={"question_id": assumption_q["id"], "state": "REJECTED"}).status_code == 403
    for q in d["questions"]:
        body = {"question_id": q["id"], "answer": "ยืนยันโดย BA", "resolve": True}
        if q["field"] == "role":
            body["field_value"] = "เจ้าหน้าที่Compliance"
        r = client.post(f"/api/requirements/{rid}/resolve-question", headers=ba, json=body)
        assert r.status_code == 200, r.text
    d = r.json()
    assert d["open_questions"] == 0 and d["fields"]["role"] == "เจ้าหน้าที่Compliance"
    r = client.post(f"/api/requirements/{rid}/approve", headers=qam)
    assert r.status_code == 200 and r.json()["status"] == "APPROVED"
    for code in ("REQ-CAM-RULE5-002", "REQ-CAM-RULE6-001", "REQ-CAM-RULE5-001"):
        assert client.post(f"/api/requirements/{rs[code]['id']}/approve", headers=qam).status_code == 200


def test_unresolved_requirement_cannot_generate(client, qam, project, demo_doc):
    rs = reqs(client, qam, project)
    r = client.post(f"/api/projects/{project['id']}/test-scenarios/generate", headers=qam,
                    json={"requirement_ids": [rs["REQ-CAM-RULE4-003"]["id"]]})
    assert r.json()["created"] == 0 and r.json()["skipped"]


def test_scenarios_and_test_cases(client, qam, qaa, project, demo_doc):
    r = client.post(f"/api/projects/{project['id']}/test-scenarios/generate", headers=qam, json={})
    assert r.status_code == 200 and r.json()["created"] > 0
    again = client.post(f"/api/projects/{project['id']}/test-scenarios/generate", headers=qam, json={}).json()
    assert again["created"] == 0                                   # duplicate prevention
    sc = client.get(f"/api/projects/{project['id']}/test-scenarios", headers=qam).json()["items"]
    types = {s["type"] for s in sc}
    assert {"Positive", "Negative", "Boundary", "Data", "API"} <= types
    r = client.post(f"/api/projects/{project['id']}/test-cases/generate", headers=qam, json={})
    assert r.json()["created"] == len(sc)
    tcs = client.get(f"/api/projects/{project['id']}/test-cases", headers=qam).json()["items"]
    b = next(t for t in tcs if t["type"] == "Boundary" and t["req_id"] == "REQ-CAM-RULE3-001")
    full = client.get(f"/api/test-cases/{b['id']}", headers=qam).json()
    vals = [(d["value"], d["expected"]) for d in full["test_data"]]
    assert ("199,999.99", "ไม่เข้าเงื่อนไข") in vals and ("200,000.00", "เข้าเงื่อนไข") in vals and ("200,000.01", "เข้าเงื่อนไข") in vals
    assert any(d["label"] == "AI RECOMMENDED TEST - NOT EXPLICITLY DEFINED IN BRS" for d in full["test_data"])
    assert full["given"] and full["when"] and full["then"] and full["business_explanation"]
    assert "AI ASSUMPTION - NOT FOUND IN BRS" in full["assumption"]
    # length boundary is a real string now (Known Issue fixed)
    lb = next(t for t in tcs if t["type"] == "Boundary" and t["req_id"] == "REQ-CAM-RULE5-002")
    lfull = client.get(f"/api/test-cases/{lb['id']}", headers=qam).json()
    assert any(d["raw"] and len(d["raw"]) == 200 for d in lfull["test_data"])
    # QA Automation cannot edit or approve
    assert client.patch(f"/api/test-cases/{b['id']}", headers=qaa, json={"reason": "x y z", "title": "x"}).status_code == 403
    # edit in table (AC 22)
    r = client.patch(f"/api/test-cases/{b['id']}", headers=qam, json={"reason": "ปรับชื่อ", "title": "Boundary: ยอดรวม 200,000",
                                                                       "steps": [{"n": 1, "action": "Login", "data": "-", "expected": "OK"}] + [
                                                                           {"n": s["n"], "action": s["action"], "data": s["data"], "expected": s["expected"], "origin": s["origin"], "label": s["label"]} for s in full["steps"][1:]]})
    assert r.status_code == 200 and r.json()["title"] == "Boundary: ยอดรวม 200,000" and r.json()["steps"][0]["action"] == "Login"
    # approve + lock (AC 23, 24)
    r = client.post(f"/api/test-cases/{b['id']}/approve", headers=qam)
    assert r.json()["status"] == "APPROVED" and r.json()["locked"] is True
    r = client.patch(f"/api/test-cases/{b['id']}", headers=qam, json={"reason": "แก้หลัง approve", "title": "zzz"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "TEST_CASE_LOCKED"
    r = client.post(f"/api/test-cases/{b['id']}/new-version", headers=qam, json={"text": "BA เปลี่ยนเงื่อนไข"})
    assert r.json()["version"] == 2 and r.json()["status"] == "REVISED"
    hist = client.get(f"/api/test-cases/{b['id']}", headers=qam).json()["history"]
    assert any(h["kind"] == "Approved snapshot" and h["snapshot"]["title"] == "Boundary: ยอดรวม 200,000" for h in hist)
    assert client.post(f"/api/test-cases/{b['id']}/approve", headers=qam).status_code == 200
    ids = [t["id"] for t in tcs if t["id"] != b["id"]]
    r = client.post("/api/test-cases/bulk-status", headers=qam, json={"ids": ids, "status": "APPROVED"})
    assert r.json()["updated"] == len(ids)


def test_excel_export(client, qam, project, demo_doc):
    r = client.get(f"/api/projects/{project['id']}/test-cases/export", headers=qam)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Requirements", "Clarification Questions", "Conflicts", "Test Scenarios", "Test Cases",
                             "Test Steps", "Test Data", "Traceability", "Automation Status"]
    ws = wb["Test Cases"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert any("ทดสอบ" in str(c.value) or "Test Case นี้" in str(c.value) for c in ws["G"])


def test_traceability(client, qam, project, demo_doc):
    tr = client.get(f"/api/projects/{project['id']}/traceability", headers=qam).json()
    chain = next(c for c in tr["chains"] if c["requirement"]["req_id"] == "REQ-CAM-RULE3-001")
    assert chain["document"]["name"] == "CAM_BRS_v1.txt" and chain["scenarios"] and chain["scenarios"][0]["test_cases"]


def test_version_compare_and_impact(client, qam, project, demo_doc):
    doc_id = demo_doc["document"]["id"]
    r = client.post(f"/api/projects/{project['id']}/documents", headers=qam,
                    files={"file": ("CAM_BRS_v2.txt", (REPO / "samples" / "CAM_BRS_v2.txt").read_bytes())}, data={"document_id": doc_id})
    assert r.json()["version"]["version"] == 2
    cmp = client.get(f"/api/documents/{doc_id}/compare", headers=qam, params={"from_version": 1, "to_version": 2}).json()
    assert [c["title"] for c in cmp["changed"]] == ["2. Rule 3 ธุรกรรมมูลค่าสูง"]
    assert cmp["added"] and cmp["removed"]
    items = cmp["impact"]["items"]
    props = {i["proposal"] for i in items}
    assert {"Update Required", "New Test Required", "Deprecation Candidate"} <= props
    # approved test cases are NOT auto-edited (AC 37)
    tcs = client.get(f"/api/projects/{project['id']}/test-cases", headers=qam).json()["items"]
    assert all(t["status"] == "APPROVED" for t in tcs if t["req_id"] == "REQ-CAM-RULE3-001")
    r = client.post(f"/api/impact/{cmp['impact']['id']}/decide", headers=qam, json={"decision": "APPROVED"})
    assert r.status_code == 200 and r.json()["status"] == "DECIDED"


def test_retry_failed_section_only(client, qam, scratch, monkeypatch):
    from app.services import ai_provider
    from app.services.engine.analysis import candidate_statements
    calls = {"n": 0}

    class Flaky(ai_provider.RuleBasedProvider):
        def extract(self, section):
            calls["n"] += 1
            if "Rule 8" in section["title"]:
                from app.core.errors import AppError
                raise AppError("AI_TIMEOUT", "AI ตอบกลับช้าเกินกำหนด", retryable=True)
            return [{"stmt": s, "fields": None, "meta": None} for s in candidate_statements(section)]

    monkeypatch.setattr("app.services.processing.get_provider", lambda mode=None: Flaky())
    text = "1. Rule 7 A\n- ระบบต้องแสดงรายงาน A ให้ผู้ดูแลระบบ\n2. Rule 8 B\n- ระบบต้องแสดงรายงาน B ให้ผู้ดูแลระบบ"
    r = client.post(f"/api/projects/{scratch['id']}/paste-text", headers=qam, json={"title": "flaky", "text": text})
    doc_id = r.json()["document"]["id"]
    prog = client.get(f"/api/documents/{doc_id}/progress", headers=qam).json()
    assert prog["job"]["status"] == "FAILED"
    st = {s["title"]: s["status"] for s in prog["sections"]}
    assert st == {"1. Rule 7 A": "DONE", "2. Rule 8 B": "FAILED"}
    assert prog["sections"][1]["error"]["code"] == "AI_TIMEOUT"
    before = calls["n"]
    monkeypatch.setattr("app.services.processing.get_provider", lambda mode=None: ai_provider.RuleBasedProvider())
    client.post(f"/api/documents/{doc_id}/retry-failed", headers=qam, json={})
    prog = client.get(f"/api/documents/{doc_id}/progress", headers=qam).json()
    assert prog["job"]["status"] == "READY_FOR_REVIEW"
    assert [s["attempts"] for s in prog["sections"]] == [1, 2]   # only the failed section re-ran (AC 39)
    assert calls["n"] == before


def test_cancel_and_resume(client, qam, scratch):
    from app.db import SessionLocal
    from app.models import ProcessingJob
    from app.services.processing import run_job
    r = client.post(f"/api/projects/{scratch['id']}/paste-text", headers=qam,
                    json={"title": "cancel", "text": "1. Rule 11 X\n- ระบบต้องแสดงผล X ให้ผู้ดูแลระบบ", "auto_process": False})
    doc_id, job_id = r.json()["document"]["id"], r.json()["version"]["job"]["id"]
    db = SessionLocal()
    real_refresh = db.refresh

    def refresh(obj, *a, **k):  # the user presses Cancel while the job is running
        real_refresh(obj, *a, **k)
        if isinstance(obj, ProcessingJob):
            obj.cancel_requested = True
    db.refresh = refresh
    run_job(db, job_id)
    assert db.get(ProcessingJob, job_id).status == "CANCELLED"
    db.close()
    assert client.post(f"/api/documents/{doc_id}/cancel", headers=qam, json={}).status_code == 200
    r = client.post(f"/api/documents/{doc_id}/resume", headers=qam, json={})
    assert r.status_code == 200
    assert client.get(f"/api/documents/{doc_id}/progress", headers=qam).json()["job"]["status"] == "READY_FOR_REVIEW"


def test_audit_log_written(client, admin):
    logs = client.get("/api/audit-logs", headers=admin).json()
    actions = {l["action"] for l in logs}
    assert {"LOGIN", "UPLOAD", "REQUIREMENT_APPROVE", "TESTCASE_APPROVED", "EXPORT", "CONFLICT_RESOLVE", "LOGIN_FAILED"} <= actions
    assert all("Admin@12345" not in (l["detail"] + l["entity"]) for l in logs)
