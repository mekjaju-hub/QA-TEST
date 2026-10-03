"""Build the Test Case catalog (Excel + CSV) from catalog_data.py + the Playwright spec files (+ last JUnit result).

    python docs/qa/build_catalog.py            # → docs/qa/WEBQA2026_TEST_CASES.xlsx + webqa2026_test_cases.csv
"""
from __future__ import annotations

import csv
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from catalog_data import CASES, DEFECTS, EXTRA, MODULES  # noqa: E402

SPECS = ROOT / "frontend" / "e2e" / "website"
JUNIT = ROOT / "storage" / "e2e-website" / "junit.xml"


def spec_titles() -> dict[str, tuple[str, str]]:
    out = {}
    for f in sorted(SPECS.glob("*.spec.ts")):
        for m in re.finditer(r'test\("(WQA-[A-Z]+-\d+) ([^"]+)"', f.read_text(encoding="utf-8")):
            out[m.group(1)] = (m.group(2), f.name)
    return out


def junit_results() -> dict[str, str]:
    if not JUNIT.exists():
        return {}
    res = {}
    for tc in ET.parse(JUNIT).getroot().iter("testcase"):
        m = re.search(r"(WQA-[A-Z]+-\d+)", tc.get("name", ""))
        if not m:
            continue
        failed = tc.find("failure") is not None or tc.find("error") is not None
        skipped = tc.find("skipped") is not None
        res[m.group(1)] = "Skipped" if skipped else ("Fail" if failed else "Pass")
    return res


def rows() -> list[dict]:
    titles, results = spec_titles(), junit_results()
    defect_of = {}
    for d in DEFECTS:
        for tid in re.findall(r"WQA-[A-Z]+-\d+", d[3]):
            defect_of.setdefault(tid, []).append(d[0])
    out = []
    for tid, (prio, typ, pre, steps, exp, data) in CASES.items():
        mod = tid.split("-")[1]
        title, fname = titles.get(tid, ("(ไม่พบใน spec)", MODULES[mod][1]))
        res = results.get(tid, "ยังไม่รัน")
        if tid == "WQA-AUTH-09" and res == "Pass":
            res = "Fail (Known defect)"           # test.fail() → Playwright reports the expected failure as passed
        out.append({"ID": tid, "Module": MODULES[mod][0], "Test Case": title, "Priority": prio, "Type": typ,
                    "Pre-condition": pre, "Steps": "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1)), "Expected Result": exp,
                    "Test Data": data, "Automation": "Playwright E2E" if mod not in ("API", "SEC", "RBAC") else "Playwright API",
                    "Script": f"frontend/e2e/website/{fname}", "Result": res, "Defect": ", ".join(defect_of.get(tid, []))})
    for tid, mod, title, prio, typ, pre, steps, exp, data, script, result, defect in EXTRA:
        out.append({"ID": tid, "Module": MODULES[mod][0], "Test Case": title, "Priority": prio, "Type": typ, "Pre-condition": pre,
                    "Steps": "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1)), "Expected Result": exp, "Test Data": data,
                    "Automation": "Manual" if script == "Manual" else ("Performance script" if script.startswith("perf/") else "Backend pytest"),
                    "Script": script, "Result": result, "Defect": defect})
    return out


HEAD = PatternFill("solid", fgColor="1F3A5F")
THIN = Side(style="thin", color="C9CED6")
FILLS = {"Pass": "D9F2E1", "ผ่าน": "D9F2E1", "Fail": "FBE0DE", "ยังไม่": "FFF3CD", "Breaking": "FFF3CD", "ฟื้นตัว": "D9F2E1"}


def style_sheet(ws, widths: list[int], wrap_cols: set[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in wrap_cols)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.freeze_panes = "B2"


def main() -> None:
    data = rows()
    cols = list(data[0].keys())
    with (HERE / "webqa2026_test_cases.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(data)

    wb = Workbook()
    s = wb.active
    s.title = "สรุป"
    s["A1"] = "WebQA2026 — Test Case ของหน้าเว็บ (Website Test Suite)"
    s["A1"].font = Font(bold=True, size=14)
    s["A2"] = f"สร้างเมื่อ {datetime.now():%d/%m/%Y %H:%M} · ผลรันล่าสุดจาก storage/e2e-website/junit.xml + perf + backend pytest"
    s.append([])
    s.append(["Module", "จำนวน TC", "อัตโนมัติ", "Manual", "ผ่าน", "ไม่ผ่าน / Known defect", "ยังไม่รัน"])
    by = {}
    for r in data:
        m = by.setdefault(r["Module"], Counter())
        m["n"] += 1
        m["auto"] += r["Automation"] != "Manual"
        m["man"] += r["Automation"] == "Manual"
        res = r["Result"]
        m["pass"] += res.startswith(("Pass", "ผ่าน", "Breaking", "ฟื้นตัว"))
        m["fail"] += res.startswith("Fail")
        m["todo"] += res.startswith("ยังไม่")
    for mod, m in by.items():
        s.append([mod, m["n"], m["auto"], m["man"], m["pass"], m["fail"], m["todo"]])
    tot = Counter()
    for m in by.values():
        tot.update(m)
    s.append(["รวม", tot["n"], tot["auto"], tot["man"], tot["pass"], tot["fail"], tot["todo"]])
    for c in s[4]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    for i, w_ in enumerate([30, 11, 11, 9, 9, 22, 11], 1):
        s.column_dimensions[get_column_letter(i)].width = w_

    ws = wb.create_sheet("Test Cases")
    ws.append(cols)
    for r in data:
        ws.append([r[c] for c in cols])
    style_sheet(ws, [14, 20, 46, 9, 12, 26, 48, 44, 18, 15, 34, 22, 14], {3, 6, 7, 8, 9, 11, 12})
    for row in ws.iter_rows(min_row=2):
        res = row[cols.index("Result")]
        for k, color in FILLS.items():
            if str(res.value).startswith(k):
                res.fill = PatternFill("solid", fgColor=color)
    t = Table(displayName="TestCases", ref=f"A1:{get_column_letter(len(cols))}{ws.max_row}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)

    wd = wb.create_sheet("Defects")
    wd.append(["Defect ID", "Severity", "Status", "พบจาก Test Case", "อาการ", "สาเหตุ", "การแก้ / ข้อเสนอ"])
    for d in DEFECTS:
        wd.append(list(d))
    style_sheet(wd, [13, 10, 18, 22, 48, 52, 48], {4, 5, 6, 7})

    wr = wb.create_sheet("วิธีรัน")
    for line in [
        ["ชุดทดสอบ", "คำสั่ง (PowerShell ที่ C:\\Cludaemek\\WebQA2026)", "หมายเหตุ"],
        ["เมนูเดียวจบ", "RUN-WEBSITE-TESTS.bat", "เลือกเมนู 1–7"],
        ["E2E + API (101 TC)", "cd frontend; npx playwright test -c playwright.website.config.ts", "เปิด backend/practice site/หน้าเว็บให้เอง (port 8002/8765/3200) · ฐานข้อมูลแยกที่ storage\\e2e-website"],
        ["เฉพาะ Module", "npx playwright test -c playwright.website.config.ts 07-web-recorder", "หรือ --grep WQA-REC-0"],
        ["ดูรายงาน", "npx playwright show-report ..\\storage\\e2e-website\\report", "มีภาพ / trace ของข้อที่ไม่ผ่าน"],
        ["Load", "python perf\\webqa_perf.py --mode load --users 50 --duration 30 --password <รหัสบัญชีทดสอบ>", "รายงาน storage\\perf\\*.html"],
        ["Stress / Spike", "--mode stress --levels 10,25,50,100,200,400  |  --mode spike --users 300", "หา Breaking point / ดูการฟื้นตัว"],
        ["Rate limit", "--mode ratelimit", "ใช้ค่า .env ปกติ (Login 5/นาที, API 600/นาที)"],
        ["Backend pytest", "cd backend; python -m pytest tests\\test_web_explorer.py", "รวม REC-14, PERF-07"],
        ["สร้างไฟล์นี้ใหม่", "python docs\\qa\\build_catalog.py", "อ่านผลจาก junit.xml ล่าสุด"],
    ]:
        wr.append(line)
    style_sheet(wr, [20, 80, 60], {2, 3})
    wb.save(HERE / "WEBQA2026_TEST_CASES.xlsx")
    print(f"{len(data)} test cases → {HERE / 'WEBQA2026_TEST_CASES.xlsx'}")


if __name__ == "__main__":
    main()
